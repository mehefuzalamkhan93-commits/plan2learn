"""Tests for the Google Drive OAuth + import integration.

Covers: OAuth state validation (unknown / expired / reused / wrong user),
callback errors, token ownership (cross-user isolation), disconnected
accounts, file selection, import failures (unsupported type, oversized,
extraction failure) and successful source creation.
"""

import asyncio
import io
from datetime import datetime, timezone

import httpx
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.ai.config import AISettings
from app.core.database import Base, get_db
from app.main import create_app
from app.models.integrations import GoogleDriveAccount, OAuthState
from app.models.notebook import Notebook, Source, SourceChunk
from app.models.user import User
from app.services import google_drive as gd


def run(coroutine):
    return asyncio.run(coroutine)


class SyncASGITestClient:
    def __init__(self, app) -> None:
        self.app = app
        self.cookies: dict[str, str] = {}

    def request(self, method: str, path: str, **kwargs) -> httpx.Response:
        async def send_request() -> httpx.Response:
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=self.app),
                base_url="http://testserver",
                cookies=self.cookies,
            ) as client:
                response = await client.request(method, path, **kwargs)
                self.cookies.update(dict(response.cookies))
                return response

        return run(send_request())

    def get(self, path: str, **kwargs) -> httpx.Response:
        return self.request("GET", path, **kwargs)

    def post(self, path: str, **kwargs) -> httpx.Response:
        return self.request("POST", path, **kwargs)


DRIVE_SETTINGS = AISettings(
    _env_file=None,
    google_drive_enabled=True,
    google_drive_client_id="test-client-id.apps.googleusercontent.com",
    google_drive_client_secret="test-client-secret",
    google_drive_redirect_uri="http://localhost:8000/api/v1/ai/drive/callback",
    google_drive_api_key="test-browser-api-key",
    google_drive_app_id="1234567890",
)


@pytest.fixture
def client_factory():
    def _make(settings: AISettings):
        engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(engine)
        testing_session = sessionmaker(bind=engine, autocommit=False, autoflush=False)

        def override_get_db():
            session = testing_session()
            try:
                yield session
            finally:
                session.close()

        app = create_app(settings)
        app.dependency_overrides[get_db] = override_get_db
        client = SyncASGITestClient(app)
        client.session_factory = testing_session
        client.settings = settings
        return client

    return _make


@pytest.fixture
def drive_client(client_factory):
    return client_factory(DRIVE_SETTINGS)


@pytest.fixture
def disabled_client(client_factory):
    return client_factory(AISettings(_env_file=None))


def signup_user(client: SyncASGITestClient, email: str = "drive_user@example.com") -> User:
    resp = client.post(
        "/api/v1/auth/signup",
        json={
            "full_name": "Drive Tester",
            "email": email,
            "password": "Password123!",
            "confirm_password": "Password123!",
        },
    )
    assert resp.status_code == 201
    with client.session_factory() as db:
        user = db.query(User).filter(User.email == email).first()
        db.expunge(user)
    return user


def create_notebook(client: SyncASGITestClient, user: User, title: str = "Drive Notebook") -> Notebook:
    with client.session_factory() as db:
        nb = Notebook(user_id=user.id, title=title)
        db.add(nb)
        db.commit()
        db.refresh(nb)
        db.expunge(nb)
    return nb


def seed_account(db_session, settings: AISettings, user: User, expired: bool = False) -> GoogleDriveAccount:
    from app.api.auth import get_current_user  # noqa: F401

    expires = datetime.now(timezone.utc).replace(
        year=2020
    ) if expired else datetime.now(timezone.utc).replace(year=2100)
    account = GoogleDriveAccount(
        user_id=user.id,
        google_email="student@gmail.com",
        access_token_encrypted=gd.encrypt_token(settings, "at-stored"),
        refresh_token_encrypted=gd.encrypt_token(settings, "rt-stored"),
        token_expiry=expires,
        scopes=gd.DRIVE_SCOPES_STRING,
        status="active",
    )
    db_session.add(account)
    db_session.commit()
    return account


# ---------------------------------------------------------------------------
# Status endpoint truthfulness
# ---------------------------------------------------------------------------

def test_status_reports_disabled_when_flag_false(disabled_client) -> None:
    signup_user(disabled_client, "status_off@example.com")
    resp = disabled_client.get("/api/v1/ai/drive/status")
    assert resp.status_code == 200
    data = resp.json()
    assert data["configured"] is False
    assert data["enabled"] is False
    assert data["connected"] is False
    assert data["connect_url"] is None


def test_status_reports_configured_but_not_connected(drive_client) -> None:
    signup_user(drive_client, "status_unconnected@example.com")
    resp = drive_client.get("/api/v1/ai/drive/status")
    data = resp.json()
    assert data["configured"] is True
    assert data["enabled"] is True
    assert data["connected"] is False
    assert data["connect_url"] == "/api/v1/ai/drive/connect"
    assert data["picker_ready"] is True
    assert set(data["scopes"]) == {
        "https://www.googleapis.com/auth/drive.file",
        "https://www.googleapis.com/auth/drive.metadata.readonly",
    }


# ---------------------------------------------------------------------------
# OAuth start + state
# ---------------------------------------------------------------------------

def test_connect_creates_single_use_user_bound_state(drive_client) -> None:
    user = signup_user(drive_client, "state_user@example.com")
    resp = drive_client.get(
        "/api/v1/ai/drive/connect", follow_redirects=False
    )
    assert resp.status_code == 307
    location = resp.headers["location"]
    assert location.startswith("https://accounts.google.com/o/oauth2/v2/auth")
    assert "state=" in location
    assert "drive.file" in location
    assert "drive.metadata.readonly" in location
    assert "access_type=offline" in location

    # State row bound to this user
    with drive_client.session_factory() as db:
        states = db.query(OAuthState).all()
        assert len(states) == 1
        assert states[0].user_id == user.id
        used_state = states[0].state


def test_connect_requires_authentication(disabled_client) -> None:
    resp = disabled_client.get("/api/v1/ai/drive/connect", follow_redirects=False)
    assert resp.status_code == 401


def test_connect_blocked_when_disabled(client_factory) -> None:
    client = client_factory(
        AISettings(
            _env_file=None,
            google_drive_client_id="id",
            google_drive_client_secret="secret",
        )
    )
    signup_user(client, "disabled_state@example.com")
    resp = client.get("/api/v1/ai/drive/connect", follow_redirects=False)
    assert resp.status_code == 403


# ---------------------------------------------------------------------------
# Callback errors
# ---------------------------------------------------------------------------

def test_callback_rejects_unknown_state(drive_client) -> None:
    signup_user(drive_client, "cb_unknown@example.com")
    resp = drive_client.get(
        "/api/v1/ai/drive/callback",
        params={"code": "some-code", "state": "not-a-real-state"},
        follow_redirects=False,
    )
    assert resp.status_code == 307
    assert "drive=error" in resp.headers["location"]
    assert "invalid_state" in resp.headers["location"]


def test_callback_rejects_reused_state(drive_client) -> None:
    signup_user(drive_client, "cb_reuse@example.com")
    drive_client.get("/api/v1/ai/drive/connect", follow_redirects=False)
    with drive_client.session_factory() as db:
        state = db.query(OAuthState).first().state

    # First (related-party-invalid) use burns the state...
    resp1 = drive_client.get(
        "/api/v1/ai/drive/callback",
        params={"code": "code-1", "state": state},
        follow_redirects=False,
    )
    assert "drive=error" in resp1.headers["location"]  # token exchange fails in tests

    # ... so replaying it must also be rejected.
    resp2 = drive_client.get(
        "/api/v1/ai/drive/callback",
        params={"code": "code-2", "state": state},
        follow_redirects=False,
    )
    assert "invalid_state" in resp2.headers["location"]


def test_callback_rejects_expired_state(drive_client) -> None:
    signup_user(drive_client, "cb_expired@example.com")
    with drive_client.session_factory() as db:
        row = OAuthState(
            state="expired-state-token",
            user_id=db.query(User).first().id,
            created_at=datetime(2020, 1, 1, tzinfo=timezone.utc),
            expires_at=datetime(2020, 1, 1, tzinfo=timezone.utc),
            used=False,
        )
        db.add(row)
        db.commit()

    resp = drive_client.get(
        "/api/v1/ai/drive/callback",
        params={"code": "code", "state": "expired-state-token"},
        follow_redirects=False,
    )
    assert "drive=error" in resp.headers["location"]
    assert "invalid_state" in resp.headers["location"]


def test_callback_error_param_propagates(drive_client) -> None:
    signup_user(drive_client, "cb_consent@example.com")
    drive_client.get("/api/v1/ai/drive/connect", follow_redirects=False)
    with drive_client.session_factory() as db:
        state = db.query(OAuthState).first().state
    resp = drive_client.get(
        "/api/v1/ai/drive/callback",
        params={"error": "access_denied", "state": state},
        follow_redirects=False,
    )
    assert resp.status_code == 307
    assert "drive=error" in resp.headers["location"]
    assert "google_access_denied" in resp.headers["location"]


# ---------------------------------------------------------------------------
# Successful OAuth callback (mocked token exchange)
# ---------------------------------------------------------------------------

def _mock_google_transport(handler) -> None:
    def fake_client() -> httpx.AsyncClient:
        return httpx.AsyncClient(
            transport=httpx.MockTransport(handler), follow_redirects=False
        )

    gd._async_client = fake_client  # type: ignore[method-assign]


def test_callback_rejects_state_from_other_session(drive_client) -> None:
    """A state minted for user A must be rejected when user B's session
    presents it in the callback (CSRF account-swap protection)."""
    user_a = signup_user(drive_client, "cb_owner@example.com")
    drive_client.get("/api/v1/ai/drive/connect", follow_redirects=False)
    with drive_client.session_factory() as db:
        state_row = db.query(OAuthState).first()
        assert state_row is not None
        assert state_row.user_id == user_a.id
        state = state_row.state

    # Fresh session logged in as user B
    client_b = SyncASGITestClient(drive_client.app)
    signup_resp = client_b.post(
        "/api/v1/auth/signup",
        json={
            "full_name": "Callback Impersonator",
            "email": "impersonator@example.com",
            "password": "Password123!",
            "confirm_password": "Password123!",
        },
    )
    assert signup_resp.status_code == 201

    resp = client_b.get(
        "/api/v1/ai/drive/callback",
        params={"code": "some-code", "state": state},
        follow_redirects=False,
    )
    assert resp.status_code == 307
    assert "drive=error" in resp.headers["location"]
    assert "invalid_state" in resp.headers["location"]

    # No account must have been created for anyone
    with drive_client.session_factory() as db:
        assert db.query(GoogleDriveAccount).count() == 0


def test_callback_success_binds_account_to_state_user(drive_client) -> None:
    user = signup_user(drive_client, "cb_success@example.com")
    drive_client.get("/api/v1/ai/drive/connect", follow_redirects=False)
    with drive_client.session_factory() as db:
        state = db.query(OAuthState).first().state

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "oauth2.googleapis.com" and "token" in str(request.url):
            assert "client_id=test-client-id" in (request.content.decode() or "")
            return httpx.Response(
                200,
                json={
                    "access_token": "at-new",
                    "refresh_token": "rt-new",
                    "expires_in": 3600,
                    "scope": gd.DRIVE_SCOPES_STRING,
                },
            )
        if request.url.host == "openidconnect.googleapis.com":
            return httpx.Response(200, json={"email": "student@gmail.com"})
        return httpx.Response(400)

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        resp = drive_client.get(
            "/api/v1/ai/drive/callback",
            params={"code": "valid-code", "state": state},
            follow_redirects=False,
        )
    finally:
        gd._async_client = original

    assert resp.status_code == 307
    assert "drive=connected" in resp.headers["location"]

    with drive_client.session_factory() as db:
        # State is burned
        assert db.query(OAuthState).filter(OAuthState.state == state).first() is None
        account = db.query(GoogleDriveAccount).first()
        assert account is not None
        assert account.user_id == user.id
        assert account.google_email == "student@gmail.com"
        assert account.status == "active"
        # Tokens stored encrypted — plaintext never persisted
        stored = db.query(GoogleDriveAccount).first()
        assert "at-new" not in (stored.access_token_encrypted or "")
        assert gd.decrypt_token(drive_client.settings, stored.access_token_encrypted) == "at-new"


# ---------------------------------------------------------------------------
# Token ownership / cross-user isolation
# ---------------------------------------------------------------------------

def test_other_user_cannot_use_connected_account(drive_client) -> None:
    user_a = signup_user(drive_client, "user_a@example.com")
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user_a)

    # Fresh session for user B
    client_b_client = SyncASGITestClient(drive_client.app)
    signup_resp = client_b_client.post(
        "/api/v1/auth/signup",
        json={
            "full_name": "Other User",
            "email": "user_b@example.com",
            "password": "Password123!",
            "confirm_password": "Password123!",
        },
    )
    assert signup_resp.status_code == 201
    # Now authenticate user B in that client (grab cookie from signup)
    # (signup set the cookie already)

    # User B must see connected=False
    status_b = client_b_client.get("/api/v1/ai/drive/status")
    assert status_b.json()["connected"] is False

    # User B cannot list files via A's connection
    files_b = client_b_client.get("/api/v1/ai/drive/files")
    assert files_b.status_code == 401

    # User B cannot disconnect A's association (it's not theirs)
    disc_b = client_b_client.post("/api/v1/ai/drive/disconnect")
    assert disc_b.status_code == 200
    with drive_client.session_factory() as db:
        remaining = (
            db.query(GoogleDriveAccount).filter(GoogleDriveAccount.user_id == user_a.id).first()
        )
        assert remaining is not None  # A's account untouched


def test_disconnect_revokes_and_removes_credentials(drive_client) -> None:
    user = signup_user(drive_client, "dc_user@example.com")
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user)

    revoked_holder: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "oauth2.googleapis.com" and "revoke" in str(request.url):
            revoked_holder.append(request.url.params.get("token", ""))
            return httpx.Response(200)
        return httpx.Response(400)

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        # Expire token first so refreshes are attempted? No — access token valid,
        # revoke uses the stored access token.
        resp = drive_client.post("/api/v1/ai/drive/disconnect")
    finally:
        gd._async_client = original

    assert resp.status_code == 200
    assert revoked_holder and revoked_holder[0] == "at-stored"

    with drive_client.session_factory() as db:
        assert (
            db.query(GoogleDriveAccount).filter(GoogleDriveAccount.user_id == user.id).first()
            is None
        )

    # Files endpoint now reports not connected
    files = drive_client.get("/api/v1/ai/drive/files")
    assert files.status_code == 401
    assert files.json()["detail"]


# ---------------------------------------------------------------------------
# File listing (metadata-only, mocked Drive API)
# ---------------------------------------------------------------------------

_DRIVE_FILES_RESPONSE = {
    "files": [
        {
            "id": "file-pdf-1",
            "name": "Lecture.pdf",
            "mimeType": "application/pdf",
            "size": "2048",
            "modifiedTime": "2026-10-01T10:00:00.000Z",
        },
        {
            "id": "file-gdoc-1",
            "name": "Notes",
            "mimeType": "application/vnd.google-apps.document",
            "modifiedTime": "2026-10-02T10:00:00.000Z",
        },
    ],
    "nextPageToken": None,
}


def test_files_listing_returns_metadata_list(drive_client) -> None:
    user = signup_user(drive_client, "files_user@example.com")
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user)

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.host == "www.googleapis.com"
        assert "Bearer at-stored" in request.headers.get("authorization", "")
        # metadata-only fields requested — no full-content scope leakage
        assert "fields" in dict(request.url.params)
        return httpx.Response(200, json=_DRIVE_FILES_RESPONSE)

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        resp = drive_client.get("/api/v1/ai/drive/files")
    finally:
        gd._async_client = original

    assert resp.status_code == 200
    data = resp.json()
    assert len(data["files"]) == 2
    names = {f["name"] for f in data["files"]}
    assert names == {"Lecture.pdf", "Notes"}


# ---------------------------------------------------------------------------
# Import: success, duplicates, and explicit failures
# ---------------------------------------------------------------------------

def test_import_google_doc_exports_and_creates_ready_source(drive_client) -> None:
    user = signup_user(drive_client, "import_ok@example.com")
    nb = create_notebook(drive_client, user)
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user)

    # Build a real DOCX payload for the exporter response
    import docx as docx_lib

    buffer = io.BytesIO()
    d = docx_lib.Document()
    d.add_heading("Photosynthesis Overview", level=1)
    d.add_paragraph("Chloroplasts convert light energy into chemical energy stored as glucose.")
    d.save(buffer)
    docx_bytes = buffer.getvalue()

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if request.url.host == "www.googleapis.com" and "/export?" in url:
            assert "openxmlformats-officedocument.wordprocessingml.document" in url
            return httpx.Response(200, content=docx_bytes)
        if request.url.host == "www.googleapis.com":
            return httpx.Response(
                200,
                json={
                    "id": "gdoc-1",
                    "name": "Photosynthesis Notes",
                    "mimeType": "application/vnd.google-apps.document",
                },
            )
        raise AssertionError(f"unexpected url {url}")

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        resp = drive_client.post(
            "/api/v1/ai/drive/import",
            json={"file_id": "gdoc-1", "notebook_id": nb.id},
        )
    finally:
        gd._async_client = original

    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["source"]["status"] == "ready"
    assert data["source"]["origin"] == "google_drive"
    assert data["chunks_indexed"] >= 1

    with drive_client.session_factory() as db:
        source = db.query(Source).first()
        assert source is not None and source.status == "ready"
        assert source.user_id == user.id
        assert source.notebook_id == nb.id
        assert source.meta_info["drive_file_id"] == "gdoc-1"
        chunks = db.query(SourceChunk).filter(SourceChunk.source_id == source.id).all()
        assert chunks
        assert "chloroplasts" in " ".join(c.content.lower() for c in chunks)


def test_import_duplicate_returns_existing_source(drive_client) -> None:
    user = signup_user(drive_client, "import_dup@example.com")
    nb = create_notebook(drive_client, user)
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user)
    docx_bytes = None
    import docx as docx_lib

    buffer = io.BytesIO()
    d = docx_lib.Document()
    d.add_paragraph("Unique duplicate-test content for the chunking pipeline.")
    d.save(buffer)
    docx_bytes = buffer.getvalue()

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "fields" in dict(request.url.params) and "/files/gdoc-dup" in url:
            return httpx.Response(
                200,
                json={"id": "gdoc-dup", "name": "Notes", "mimeType": "application/vnd.google-apps.document"},
            )
        if "/export?" in url:
            return httpx.Response(200, content=docx_bytes)
        raise AssertionError(url)

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        first = drive_client.post("/api/v1/ai/drive/import", json={"file_id": "gdoc-dup", "notebook_id": nb.id})
        assert first.status_code == 200
        second = drive_client.post("/api/v1/ai/drive/import", json={"file_id": "gdoc-dup", "notebook_id": nb.id})
    finally:
        gd._async_client = original

    assert second.status_code == 200
    data = second.json()
    assert data["duplicate"] is True
    assert data["source"]["id"] == first.json()["source"]["id"]

    with drive_client.session_factory() as db:
        assert db.query(Source).count() == 1


def test_import_rejects_unsupported_file_type(drive_client) -> None:
    user = signup_user(drive_client, "import_sheet@example.com")
    nb = create_notebook(drive_client, user)
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user)

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "id": "sheet-1",
                "name": "Grades",
                "mimeType": "application/vnd.google-apps.spreadsheet",
            },
        )

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        resp = drive_client.post("/api/v1/ai/drive/import", json={"file_id": "sheet-1", "notebook_id": nb.id})
    finally:
        gd._async_client = original

    assert resp.status_code == 422
    assert "not supported" in resp.json()["detail"]


def test_import_rejects_oversized_file(drive_client) -> None:
    user = signup_user(drive_client, "import_big@example.com")
    nb = create_notebook(drive_client, user)
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user)

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "id": "pdf-big",
                "name": "Huge.pdf",
                "mimeType": "application/pdf",
                "size": str(60 * 1024 * 1024),
            },
        )

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        resp = drive_client.post("/api/v1/ai/drive/import", json={"file_id": "pdf-big", "notebook_id": nb.id})
    finally:
        gd._async_client = original

    assert resp.status_code == 422
    assert "50 MB" in resp.json()["detail"]


def test_import_extraction_failure_marks_source_failed(drive_client) -> None:
    user = signup_user(drive_client, "import_fail@example.com")
    nb = create_notebook(drive_client, user)
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user)

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "fields" in dict(request.url.params):
            return httpx.Response(
                200,
                json={"id": "pdf-broken", "name": "corrupted.pdf", "mimeType": "application/pdf"},
            )
        # Return garbage that isn't a valid PDF; extractor must fail.
        return httpx.Response(200, content=b"%PDF-broken-no-text-here")

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        resp = drive_client.post("/api/v1/ai/drive/import", json={"file_id": "pdf-broken", "notebook_id": nb.id})
    finally:
        gd._async_client = original

    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail  # honest error, never a fake success

    with drive_client.session_factory() as db:
        source = db.query(Source).filter(Source.title == "corrupted.pdf").first()
        assert source is not None
        assert source.status == "failed"
        assert source.error_message
        assert db.query(SourceChunk).filter(SourceChunk.source_id == source.id).count() == 0


def test_import_requires_notebook_ownership(drive_client) -> None:
    user = signup_user(drive_client, "import_own@example.com")
    nb = create_notebook(drive_client, user)

    # A different user attempts to import into user's notebook
    other_client = SyncASGITestClient(drive_client.app)
    signup_resp = other_client.post(
        "/api/v1/auth/signup",
        json={
            "full_name": "Import Intruder",
            "email": "intruder@example.com",
            "password": "Password123!",
            "confirm_password": "Password123!",
        },
    )
    assert signup_resp.status_code == 201

    resp = other_client.post(
        "/api/v1/ai/drive/import", json={"file_id": "file-x", "notebook_id": nb.id}
    )
    assert resp.status_code == 404  # notebook not visible to the other user


def test_import_requires_connection(drive_client) -> None:
    user = signup_user(drive_client, "import_noconn@example.com")
    nb = create_notebook(drive_client, user)
    resp = drive_client.post("/api/v1/ai/drive/import", json={"file_id": "file-1", "notebook_id": nb.id})
    assert resp.status_code == 401


# ---------------------------------------------------------------------------
# Picker
# ---------------------------------------------------------------------------

def test_picker_config_requires_connection_and_configuration(drive_client) -> None:
    user = signup_user(drive_client, "picker@example.com")
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user)

    cfg = drive_client.post("/api/v1/ai/drive/picker")
    assert cfg.status_code == 200
    data = cfg.json()
    assert data["client_id"] == DRIVE_SETTINGS.google_drive_client_id
    assert data["app_id"] == DRIVE_SETTINGS.google_drive_app_id

    tok = drive_client.post("/api/v1/ai/drive/picker/token")
    assert tok.status_code == 200
    assert tok.json()["access_token"] == "at-stored"

    # Client without a connection may not obtain picker tokens
    noconn = SyncASGITestClient(drive_client.app)
    signup_resp = noconn.post(
        "/api/v1/auth/signup",
        json={
            "full_name": "Picker Stranger",
            "email": "picker_stranger@example.com",
            "password": "Password123!",
            "confirm_password": "Password123!",
        },
    )
    assert signup_resp.status_code == 201
    assert noconn.post("/api/v1/ai/drive/picker/token").status_code == 401


def test_expired_access_token_triggers_refresh(drive_client) -> None:
    user = signup_user(drive_client, "refresh@example.com")
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user, expired=True)

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if request.url.host == "oauth2.googleapis.com" and "token" in url:
            body = request.content.decode()
            assert "grant_type=refresh_token" in body
            assert "refresh_token=rt-stored" in body
            return httpx.Response(
                200,
                json={"access_token": "at-refreshed", "expires_in": 3600},
            )
        if request.url.host == "www.googleapis.com":
            assert request.headers.get("authorization") == "Bearer at-refreshed"
            return httpx.Response(200, json={"files": [], "nextPageToken": None})
        raise AssertionError(url)

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        resp = drive_client.get("/api/v1/ai/drive/files")
    finally:
        gd._async_client = original

    assert resp.status_code == 200
    with drive_client.session_factory() as db:
        account = db.query(GoogleDriveAccount).first()
        assert gd.decrypt_token(drive_client.settings, account.access_token_encrypted) == "at-refreshed"


def test_invalid_grant_marks_account_needs_reauth(drive_client) -> None:
    user = signup_user(drive_client, "reauth@example.com")
    with drive_client.session_factory() as db:
        seed_account(db, drive_client.settings, user, expired=True)

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "oauth2.googleapis.com":
            return httpx.Response(400, json={"error": "invalid_grant"})
        raise AssertionError(str(request.url))

    original = gd._async_client
    _mock_google_transport(handler)
    try:
        resp = drive_client.get("/api/v1/ai/drive/files")
    finally:
        gd._async_client = original

    assert resp.status_code == 401
    assert "reconnect" in resp.json()["detail"].lower()
    with drive_client.session_factory() as db:
        account = db.query(GoogleDriveAccount).first()
        assert account.status == "needs_reauth"
